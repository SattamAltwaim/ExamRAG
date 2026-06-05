import os
import re
import time
from pathlib import Path

from pptx import Presentation
import voyageai
import chromadb

import config


def extract_slide_content(pptx_path: str) -> list[dict]:
    filename = Path(pptx_path).name
    topic = re.sub(r"^\d+\.?\s*", "", Path(pptx_path).stem).strip()
    prs = Presentation(pptx_path)
    slides = []

    for i, slide in enumerate(prs.slides, start=1):
        slide_text_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_text_parts.append(text)

        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()

        slide_text = "\n".join(slide_text_parts)
        title = slide_text_parts[0] if slide_text_parts else ""
        combined = f"{slide_text}\n{notes_text}".strip()

        if len(combined) < 20:
            continue

        slides.append({
            "source_file": filename,
            "slide_number": i,
            "slide_title": title[:200],
            "topic": topic,
            "combined_text": combined,
        })

    return slides


def ingest_all_content():
    content_dir = Path(config.CONTENT_DIR)
    all_chunks = []

    for f in sorted(content_dir.iterdir()):
        if f.suffix.lower() in (".pptx", ".pptm"):
            print(f"Extracting: {f.name}")
            chunks = extract_slide_content(str(f))
            all_chunks.extend(chunks)
            print(f"  → {len(chunks)} slides extracted")

    if not all_chunks:
        print("No content found.")
        return

    print(f"\nTotal chunks: {len(all_chunks)}")
    print("Embedding with Voyage 3...")

    vo = voyageai.Client(api_key=config.VOYAGE_API_KEY)
    texts = [c["combined_text"] for c in all_chunks]

    batch_size = 20
    all_embeddings = []
    for start in range(0, len(texts), batch_size):
        batch = texts[start : start + batch_size]
        result = vo.embed(batch, model=config.EMBEDDING_MODEL, input_type="document")
        all_embeddings.extend(result.embeddings)
        done = min(start + batch_size, len(texts))
        print(f"  Embedded {done}/{len(texts)}")
        if done < len(texts):
            time.sleep(22)

    print("Storing in ChromaDB...")
    client = chromadb.PersistentClient(path=config.CHROMA_DB_PATH)
    collection = client.get_or_create_collection(
        name=config.COLLECTION_NAME,
        metadata={"hnsw:space": "cosine"},
    )

    ids = [f"{c['source_file']}__slide_{c['slide_number']}" for c in all_chunks]
    metadatas = [
        {
            "source_file": c["source_file"],
            "topic": c["topic"],
            "slide_number": c["slide_number"],
            "slide_title": c["slide_title"],
        }
        for c in all_chunks
    ]

    collection.upsert(
        ids=ids,
        embeddings=all_embeddings,
        documents=texts,
        metadatas=metadatas,
    )

    print(f"Done. {collection.count()} documents in collection.")


if __name__ == "__main__":
    ingest_all_content()
